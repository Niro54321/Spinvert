"""Run every operation on generated sample files and check the outputs.

    python tests/smoke_test.py [work-folder] [--no-office]

Needs FFmpeg on PATH. Uses Word, PowerPoint and Excel when installed, unless
--no-office is given (that also tests the Qt PDF fallback).
"""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PIL import Image, ImageDraw  # noqa: E402

from spinvert.catalog import choices_for, detect_engines  # noqa: E402
from spinvert.core import Context, run_op  # noqa: E402
from spinvert.settings import Settings  # noqa: E402
from spinvert.util import ConversionError, find_ffmpeg, find_ffprobe  # noqa: E402

MARKDOWN = """# Sample document

Some **bold** text, some *italic* text, and an image:

![chart](chart.png)

- first point
- second point

| name | value |
| ---- | ----- |
| a    | 1     |
| b    | 2     |
"""


def make_samples(folder: Path, ffmpeg: Path) -> dict[str, list[Path]]:
    s = {}
    img = Image.new("RGBA", (640, 400), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    for i in range(0, 640, 8):
        d.line([(i, 0), (i, 400)], fill=(i % 256, 120, 255 - i % 256, 255), width=8)
    d.ellipse([200, 100, 440, 300], fill=(255, 255, 255, 0))
    png = folder / "photo.png"
    img.save(png)
    img.convert("RGB").save(folder / "chart.png")
    jpg = folder / "camera.jpg"
    exif = Image.Exif()
    exif[0x0112] = 6  # rotated: tests the upright fix
    img.convert("RGB").save(jpg, quality=95, exif=exif.tobytes())
    frames = [Image.new("RGB", (120, 80), c) for c in ("red", "green", "blue")]
    gif = folder / "anim.gif"
    frames[0].save(gif, save_all=True, append_images=frames[1:], duration=200, loop=0)
    webp = folder / "still.webp"
    img.save(webp, quality=90)
    tiff = folder / "scan.tiff"
    img.convert("RGB").save(tiff)
    bmp = folder / "old.bmp"
    img.convert("RGB").save(bmp)
    s["image"] = [png, jpg, gif, webp, tiff, bmp]
    try:
        import pillow_heif
        heic = folder / "iphone.heic"
        pillow_heif.from_pillow(img.convert("RGB")).save(str(heic), quality=80)
        s["image"].append(heic)
    except Exception as e:
        print("  (no HEIC sample:", e, ")")
    try:
        avif = folder / "modern.avif"
        img.convert("RGB").save(avif, quality=70)
        s["image"].append(avif)
    except Exception as e:
        print("  (no AVIF sample:", e, ")")

    def ff(*args):
        subprocess.run([str(ffmpeg), "-hide_banner", "-loglevel", "error", "-y", *args], check=True)

    wav = folder / "tone.wav"
    ff("-f", "lavfi", "-i", "sine=frequency=440:duration=3", str(wav))
    mp3 = folder / "song.mp3"
    ff("-i", str(wav), "-c:a", "libmp3lame", str(mp3))
    s["audio"] = [wav, mp3]
    mp4 = folder / "clip.mp4"
    ff("-f", "lavfi", "-i", "testsrc=duration=3:size=640x360:rate=25", "-f", "lavfi",
       "-i", "sine=duration=3", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac",
       "-shortest", str(mp4))
    avi = folder / "home.avi"
    ff("-i", str(mp4), "-c:v", "mpeg4", "-c:a", "libmp3lame", str(avi))
    s["video"] = [mp4, avi]

    pages = [Image.new("RGB", (1240, 1754), c) for c in ("white", "lightyellow", "lightblue")]
    for n, p in enumerate(pages, 1):
        ImageDraw.Draw(p).text((100, 100), f"Page {n}", fill="black")
    big = Image.merge("RGB", [Image.effect_mandelbrot((2400, 1800), (-2, -1.2, 1, 1.2), 60),
                              Image.effect_noise((2400, 1800), 60),
                              Image.linear_gradient("L").resize((2400, 1800))])
    pdf = folder / "report.pdf"
    pages[0].save(pdf, save_all=True, append_images=pages[1:] + [big], resolution=150)
    pdf2 = folder / "appendix.pdf"
    pages[1].save(pdf2, resolution=150)
    from spinvert.qtpdf import html_to_pdf
    text_pdf = folder / "article.pdf"
    html_to_pdf("<h1>Article</h1>" + "<p>Real selectable text in a PDF.</p>" * 40, folder, text_pdf)
    s["pdf"] = [pdf, pdf2, text_pdf]

    md = folder / "notes.md"
    md.write_text(MARKDOWN, encoding="utf-8")
    from spinvert.engines.documents import pandoc_path
    pandoc = pandoc_path()
    docx = folder / "letter.docx"
    subprocess.run([str(pandoc), md.name, "-o", docx.name], cwd=folder, check=True)
    epub = folder / "book.epub"
    subprocess.run([str(pandoc), md.name, "-M", "title=Book", "-o", epub.name], cwd=folder, check=True)
    html = folder / "page.html"
    html.write_text("<html><body><h1>Hello</h1><p>A <b>web</b> page.</p></body></html>", encoding="utf-8")
    txt = folder / "readme.txt"
    txt.write_text("Line one\nLine two\n\nNew paragraph with ümlaut and عربي.", encoding="utf-8")
    s["document"] = [md, docx, epub, html, txt]
    rtfd = folder / "note.rtfd"
    rtfd.mkdir()
    subprocess.run([str(pandoc), md.name, "-s", "-t", "rtf", "-o", str(rtfd / "TXT.rtf")], cwd=folder, check=True)
    clip = folder / "snippet.textClipping"
    import plistlib
    clip.write_bytes(plistlib.dumps({"UTIData": {"public.utf8-plain-text": "Clipped text\nsecond line".encode()}},
                                    fmt=plistlib.FMT_BINARY))
    s["document"] += [rtfd, clip]
    svg = folder / "logo.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80" viewBox="0 0 120 80">'
                   '<rect x="5" y="5" width="110" height="70" rx="14" fill="#7C5CFF"/>'
                   '<circle cx="60" cy="40" r="22" fill="#fff" opacity=".8"/></svg>', encoding="utf-8")
    s["image"].append(svg)
    pages_file = folder / "essay.pages"
    with zipfile.ZipFile(pages_file, "w") as zf:
        zf.writestr("Index/Document.iwa", b"\x00private")
        zf.write(text_pdf, "QuickLook/Preview.pdf")
    s["pages"] = [pages_file]

    from pptx import Presentation
    prs = Presentation()
    for n in (1, 2):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"Slide title {n}"
        body = slide.placeholders[1].text_frame
        body.text = "Main point"
        p = body.add_paragraph()
        p.text = "Sub point"
        p.level = 1
    pptx = folder / "deck.pptx"
    prs.save(str(pptx))
    s["slides"] = [pptx]

    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["month", "amount"])
    ws.append(["Jan", 1200.5])
    ws.append(["Feb", 980])
    wb.create_sheet("Notes").append(["note", "zip code"])
    xlsx = folder / "budget.xlsx"
    wb.save(str(xlsx))
    csv_ = folder / "people.csv"
    csv_.write_text("name;age;zip\nAna;31;01234\nOmar;45;90210\n", encoding="utf-8")
    tsv = folder / "table.tsv"
    tsv.write_text("a\tb\n1\t2\n", encoding="utf-8")
    from spinvert.engines.sheets import write_ods
    ods = folder / "open.ods"
    write_ods([("Data", [["item", "qty"], ["pen", 3], ["ink", 2.5]])], ods)
    s["sheet"] = [xlsx, csv_, tsv, ods]

    srt = folder / "movie.srt"
    srt.write_text("1\n00:00:01,000 --> 00:00:02,500\n<i>Hello</i>\n\n2\n00:00:03,000 --> 00:00:04,000\nWorld\n",
                   encoding="utf-8")
    vtt = folder / "talk.vtt"
    vtt.write_text("WEBVTT\n\nNOTE comment\n\n00:01.000 --> 00:02.000 align:start\nHi there\n",
                   encoding="utf-8")
    s["subtitle"] = [srt, vtt]

    pack = folder / "bundle"
    pack.mkdir()
    (pack / "a.txt").write_text("a", encoding="utf-8")
    (pack / "sub").mkdir()
    (pack / "sub" / "b.txt").write_text("b", encoding="utf-8")
    zip_ = folder / "bundle.zip"
    with zipfile.ZipFile(zip_, "w") as zf:
        zf.write(pack / "a.txt", "bundle/a.txt")
        zf.write(pack / "sub" / "b.txt", "bundle/sub/b.txt")
    import tarfile, gzip, shutil
    tgz = folder / "backup.tar.gz"
    with tarfile.open(tgz, "w:gz") as tf:
        tf.add(pack, "bundle")
    gz = folder / "log.txt.gz"
    with open(txt, "rb") as fin, gzip.open(gz, "wb") as fout:
        shutil.copyfileobj(fin, fout)
    s["archive"] = [zip_, tgz, gz]
    rar_tool = Path(r"C:\Program Files\WinRAR\Rar.exe")
    if rar_tool.is_file():  # only to make a sample; Spinvert itself opens RAR with Windows' tar
        rar = folder / "photos.rar"
        subprocess.run([str(rar_tool), "a", "-ep1", "-idq", str(rar), str(pack)], check=True)
        s["archive"].append(rar)
    else:
        print("  (no RAR sample: WinRAR not installed)")
    s["folder"] = [pack]
    return s


def check_output(path: Path, ffprobe: Path | None) -> str:
    if path.is_dir():
        files = list(path.rglob("*"))
        return f"folder, {len(files)} items" if files else "EMPTY FOLDER"
    if not path.exists() or path.stat().st_size == 0:
        raise AssertionError("missing or empty")
    ext = path.suffix.lower()
    if ext in (".jpg", ".png", ".webp", ".tiff", ".bmp", ".gif", ".heic", ".avif"):
        with Image.open(path) as im:
            im.load()
            return f"{im.format} {im.size[0]}x{im.size[1]} {im.mode}"
    if ext == ".pdf":
        from pypdf import PdfReader
        return f"PDF {len(PdfReader(str(path)).pages)} pages"
    if ext in (".mp3", ".m4a", ".wav", ".flac", ".ogg", ".opus", ".aiff", ".wma",
               ".mp4", ".mov", ".mkv", ".webm", ".avi", ".wmv"):
        out = subprocess.run([str(ffprobe), "-v", "error", "-show_entries",
                              "format=duration:stream=codec_name", "-of", "json", str(path)],
                             capture_output=True, text=True, check=True).stdout
        data = json.loads(out)
        codecs = ",".join(st["codec_name"] for st in data.get("streams", []))
        return f"{float(data['format']['duration']):.1f}s {codecs}"
    if ext == ".ods":
        from spinvert.engines.sheets import read_sheets
        sheets = read_sheets(None, path)
        return f"ODS {len(sheets)} sheets, {sum(len(r) for _, r in sheets)} rows"
    if ext == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(str(path), read_only=True)
        rows = sum(ws.max_row or 0 for ws in book.worksheets)
        return f"XLSX {len(book.worksheets)} sheets, {rows} rows"
    if ext in (".docx", ".pptx", ".epub", ".zip"):
        assert zipfile.is_zipfile(path), "not a valid zip container"
        return f"{ext[1:].upper()} {path.stat().st_size // 1024 + 1} KB"
    if ext in (".tar", ".gz"):
        import tarfile
        with tarfile.open(path) as tf:
            return f"TAR {len(tf.getnames())} entries"
    text = path.read_text(encoding="utf-8-sig")
    assert text.strip(), "empty text"
    return f"{len(text)} chars: {text.strip().splitlines()[0][:40]!r}"


def main() -> int:
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    no_office = "--no-office" in sys.argv
    work = Path(args[0]) if args else Path(tempfile.mkdtemp(prefix="spinvert-test-"))
    work.mkdir(parents=True, exist_ok=True)
    samples_dir = work / f"run-{int(time.time())}"
    samples_dir.mkdir()

    from PySide6.QtGui import QGuiApplication
    # The real Windows platform: offscreen has no fonts, so PDFs would get no text.
    app = QGuiApplication.instance() or QGuiApplication(["test"])  # noqa: F841

    settings = Settings()
    engines = detect_engines(settings)
    if no_office:
        engines.update(word=False, powerpoint=False, excel=False)
    print("engines:", engines)
    ffmpeg = find_ffmpeg()
    ffprobe = find_ffprobe(ffmpeg)
    samples = make_samples(samples_dir, ffmpeg)

    jobs = []
    for cat, files in samples.items():
        for f in files:
            for wheel in ("convert", "tools"):
                for choice in choices_for([f], wheel, engines):
                    jobs.append((f.name, choice))
    groups = [samples["image"][:3], samples["pdf"], [samples["image"][0], samples["pdf"][0]],
              samples["audio"] + samples["video"]]
    for group in groups:
        for wheel in ("convert", "tools"):
            for choice in choices_for(group, wheel, engines):
                if len(choice.plan) == 1 and len(choice.plan[0][1]) > 1 or wheel == "convert":
                    jobs.append((" + ".join(p.name for p in group), choice))

    failures = 0
    for name, choice in jobs:
        for op, inputs in choice.plan:
            params = {}
            if op.prepare:
                from spinvert.engines.pdfs import page_count
                params = {"order": list(reversed(range(page_count(inputs[0]))))}
            ctx = Context(inputs=inputs, settings=settings, engines=engines, params=params)
            started = time.perf_counter()
            label = f"{name[:38]:38} {choice.wheel:7} {choice.label:16}"
            try:
                outputs = run_op(op, ctx)
                info = "; ".join(check_output(o, ffprobe) for o in outputs)
                print(f"OK    {label} {time.perf_counter() - started:5.1f}s  {info}")
            except ConversionError as e:
                expected = any(s in str(e) for s in ("Already compact", "only one page", "No text found"))
                tag = "NOTE " if expected else "FAIL "
                failures += 0 if expected else 1
                print(f"{tag} {label} {e}")
            except AssertionError as e:
                failures += 1
                print(f"FAIL  {label} bad output: {e}")
            break  # one file per choice is enough for mixed groups
    print(f"\n{len(jobs)} choices tested, {failures} failures. Files in {samples_dir}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
