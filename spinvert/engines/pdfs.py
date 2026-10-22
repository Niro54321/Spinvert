"""PDF: render pages with PDFium, edit pages with pypdf."""

from __future__ import annotations

import base64
import html
import io
import re
import statistics
import tempfile
from pathlib import Path

import pypdfium2 as pdfium
from pypdf import PdfReader, PdfWriter

from spinvert.util import ConversionError, natural_key

MAX_PIXELS = 40_000_000


def _open_pdfium(path: Path):
    try:
        return pdfium.PdfDocument(str(path))
    except pdfium.PdfiumError as e:
        if "password" in str(e).lower():
            raise ConversionError(f"{path.name} is password protected.") from e
        raise ConversionError(f"Can't read {path.name}: {e}") from e


def _reader(path: Path) -> PdfReader:
    try:
        reader = PdfReader(str(path))
    except Exception as e:
        raise ConversionError(f"Can't read {path.name}: {e}") from e
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ConversionError(f"{path.name} is password protected.")
        except ConversionError:
            raise
        except Exception as e:
            raise ConversionError(f"{path.name} is password protected.") from e
    return reader


def _render(page, dpi: int):
    width, height = page.get_size()
    scale = dpi / 72
    pixels = width * scale * height * scale
    if pixels > MAX_PIXELS:
        scale *= (MAX_PIXELS / pixels) ** 0.5
    return page.render(scale=scale).to_pil()


def pdf_to_images(fmt: str):
    def run(ctx):
        src = ctx.inputs[0]
        pdf = _open_pdfium(src)
        try:
            count = len(pdf)
            if count == 1:
                targets = [ctx.output(src.stem, fmt)]
                result = targets
            else:
                folder = ctx.output_dir(f"{src.stem} pages")
                digits = len(str(count))
                targets = [folder / f"{src.stem}-p{i + 1:0{digits}d}.{fmt}" for i in range(count)]
                result = [folder]
            for i, target in enumerate(targets):
                ctx.check()
                image = _render(pdf[i], 300)
                if fmt == "jpg":
                    image.convert("RGB").save(target, "JPEG", quality=90, optimize=True)
                else:
                    image.save(target, "PNG")
                ctx.progress((i + 1) / count)
        finally:
            pdf.close()
        return result
    return run


NO_TEXT = "No text found. This PDF is probably scanned images."


def _page_texts(ctx, src: Path) -> list[str]:
    pdf = _open_pdfium(src)
    parts = []
    try:
        for i in range(len(pdf)):
            ctx.check()
            textpage = pdf[i].get_textpage()
            parts.append(textpage.get_text_bounded().replace("\r\n", "\n").strip())
            ctx.progress(0.5 * (i + 1) / len(pdf))
    finally:
        pdf.close()
    return parts


def pdf_to_text(ctx):
    src = ctx.inputs[0]
    text = "\n\n".join(_page_texts(ctx, src)).strip()
    if not text:
        raise ConversionError(NO_TEXT)
    out = ctx.output(src.stem, "txt")
    out.write_text(text + "\n", encoding="utf-8")
    return [out]


# PDF to DOCX, EPUB and Markdown. A PDF stores characters at positions, not
# paragraphs, so the document is rebuilt: lines are grouped into paragraphs by
# their spacing, bigger text becomes headings, ruled tables become tables, and
# pictures are cut out of the page. A page with no text at all (a scan) comes
# through as a picture of the whole page.

def _data_uri(image) -> str:
    buf = io.BytesIO()
    image.convert("RGB").save(buf, "JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


def _inside(obj, boxes) -> bool:
    return any(x0 - 1 <= obj["x0"] and obj["x1"] <= x1 + 1 and top - 1 <= obj["top"] and obj["bottom"] <= bottom + 1
               for x0, top, x1, bottom in boxes)


def _page_blocks(page, fitz_page, body_size: float) -> list[tuple[float, str]]:
    blocks = []
    tables = page.find_tables()
    boxes = [t.bbox for t in tables]
    for table in tables:
        rows = "".join("<tr>" + "".join(f"<td>{html.escape((c or '').strip())}</td>" for c in row) + "</tr>"
                       for row in table.extract())
        blocks.append((table.bbox[1], f"<table>{rows}</table>"))

    for img in page.images:
        if img["x1"] - img["x0"] < 30 or img["bottom"] - img["top"] < 30:
            continue
        crop = (max(0, img["x0"]), max(0, page.height - img["bottom"]),
                max(0, page.width - img["x1"]), max(0, img["top"]))
        picture = fitz_page.render(scale=150 / 72, crop=crop).to_pil()
        blocks.append((img["top"], f'<p><img src="{_data_uri(picture)}"/></p>'))

    lines = [ln for ln in page.extract_text_lines(return_chars=True) if not _inside(ln, boxes)]
    paragraph, last, last_size, last_tag = [], None, 0.0, "p"

    def flush():
        if paragraph:
            text = " ".join(t for t in paragraph)
            blocks.append((start_top, f"<{last_tag}>{html.escape(text)}</{last_tag}>"))
            paragraph.clear()

    start_top = 0.0
    for line in lines:
        chars = [c for c in line["chars"] if c["text"].strip()]
        if not chars:
            continue
        size = sum(c["size"] for c in chars) / len(chars)
        tag = "h1" if size >= body_size * 1.8 else "h2" if size >= body_size * 1.3 else "p"
        gap = line["top"] - last["bottom"] if last else 0
        if last is None or tag != last_tag or abs(size - last_size) > 1 or gap > 0.8 * size:
            flush()
            start_top = line["top"]
        paragraph.append(line["text"].strip())
        last, last_size, last_tag = line, size, tag
    flush()

    if not lines and not tables:  # a scanned page: keep it as a picture
        blocks.append((0, f'<p><img src="{_data_uri(fitz_page.render(scale=150 / 72).to_pil())}"/></p>'))
    return sorted(blocks, key=lambda b: b[0])


def _reflow_html(ctx, src: Path) -> str:
    import pdfplumber

    try:
        plumb = pdfplumber.open(str(src))
    except Exception as e:
        raise ConversionError(f"Can't read {src.name}. Is it password protected?") from e
    fitz = _open_pdfium(src)
    parts = []
    try:
        with plumb:
            sample = [c["size"] for page in plumb.pages[:20] for c in page.chars if c["text"].strip()]
            body_size = statistics.median(sample) if sample else 11.0
            count = len(plumb.pages)
            for i, page in enumerate(plumb.pages):
                ctx.check()
                parts += [block for _, block in _page_blocks(page, fitz[i], body_size)]
                ctx.progress(0.7 * (i + 1) / count)
    finally:
        fitz.close()
    title = html.escape(src.stem)
    return f"<html><head><title>{title}</title></head><body>{''.join(parts)}</body></html>"


def pdf_to(ext: str):
    """PDF to DOCX, EPUB or Markdown through the rebuilt document."""
    def run(ctx):
        from spinvert.engines.documents import WRITERS, pandoc_path

        src = ctx.inputs[0]
        out = ctx.output(src.stem, ext)
        extra = []
        if ext == "md":
            media = out.parent / f"{out.stem} files"  # the pictures, next to the Markdown file
            ctx.created.append(media)
            extra = ["--extract-media", media.name]
        elif ext == "epub":
            extra = ["--metadata", f"title={src.stem}"]
        with tempfile.TemporaryDirectory(prefix="spinvert-") as tmp:
            page = Path(tmp) / "document.html"
            page.write_text(_reflow_html(ctx, src), encoding="utf-8")
            ctx.run([pandoc_path(), page, "-f", "html", *WRITERS[ext], *extra, "-o", out.name],
                    cwd=out.parent)
        if ext == "md" and media.exists():
            return [out, media]
        return [out]
    return run


def pdf_to_xlsx(ctx):
    """Tables become rows and columns, one sheet per page that has tables. A PDF
    without ruled tables gets its text lines, split into columns at wide gaps."""
    import pdfplumber
    from openpyxl import Workbook

    from spinvert.engines.sheets import parse_cell

    src = ctx.inputs[0]
    book = Workbook()
    book.remove(book.active)
    try:
        pdf = pdfplumber.open(str(src))
    except Exception as e:
        raise ConversionError(f"Can't read {src.name}. Is it password protected?") from e
    with pdf:
        count = len(pdf.pages)
        for i, page in enumerate(pdf.pages, 1):
            ctx.check()
            tables = page.extract_tables()
            if tables:
                ws = book.create_sheet(f"Page {i}")
                for n, table in enumerate(tables):
                    if n:
                        ws.append([])
                    for row in table:
                        ws.append([parse_cell(c.strip()) if isinstance(c, str) else c for c in row])
            ctx.progress(0.9 * i / count)
        if not book.worksheets:
            ws = book.create_sheet("Text")
            for page in pdf.pages:
                for line in (page.extract_text(layout=True) or "").splitlines():
                    if line.strip():
                        ws.append([parse_cell(c) for c in re.split(r"\s{2,}", line.strip())])
            if ws.max_row <= 1 and ws.cell(1, 1).value is None:
                raise ConversionError(NO_TEXT)
    out = ctx.output(src.stem, "xlsx")
    book.save(str(out))
    return [out]


def pdf_to_pptx(ctx):
    from pptx import Presentation
    from pptx.util import Emu

    src = ctx.inputs[0]
    pdf = _open_pdfium(src)
    try:
        w_pt, h_pt = pdf[0].get_size()
        shrink = min(1.0, 4000 / max(w_pt, h_pt))
        prs = Presentation()
        prs.slide_width = Emu(int(w_pt * shrink * 12700))
        prs.slide_height = Emu(int(h_pt * shrink * 12700))
        blank = prs.slide_layouts[6]
        for i in range(len(pdf)):
            ctx.check()
            stream = io.BytesIO()
            _render(pdf[i], 150).convert("RGB").save(stream, "JPEG", quality=88)
            stream.seek(0)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(stream, 0, 0, prs.slide_width, prs.slide_height)
            ctx.progress(0.95 * (i + 1) / len(pdf))
    finally:
        pdf.close()
    out = ctx.output(src.stem, "pptx")
    prs.save(str(out))
    return [out]


def merge(ctx):
    order = ctx.params.get("order")
    files = [ctx.inputs[i] for i in order] if order else sorted(ctx.inputs, key=natural_key)
    writer = PdfWriter()
    for i, path in enumerate(files):
        ctx.check()
        writer.append(_reader(path))
        ctx.progress(0.9 * (i + 1) / len(files))
    out = ctx.output("Merged", "pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    return [out]


def parse_ranges(text: str, count: int) -> list[list[int]]:
    """'1-3, 5, 8-' to [[0,1,2],[4],[7..count-1]] (page indexes)."""
    groups = []
    for part in re.split(r"[,;\s]+", text.strip()):
        if not part:
            continue
        a, dash, b = part.partition("-")
        try:
            first = int(a) if a else 1
            last = (int(b) if b else count) if dash else first
        except ValueError as e:
            raise ConversionError(f"Can't read the page range '{part}'.") from e
        first, last = max(1, first), min(count, last)
        if first <= last:
            groups.append(list(range(first - 1, last)))
    if not groups:
        raise ConversionError("No pages in those ranges.")
    return groups


def split(ctx):
    """Every page, every N pages, or custom ranges like '1-3, 4-7, 8-'."""
    src = ctx.inputs[0]
    reader = _reader(src)
    count = len(reader.pages)
    p = ctx.params
    mode = p.get("mode", "each")
    if mode == "every":
        n = max(1, int(p.get("every", 2)))
        groups = [list(range(i, min(count, i + n))) for i in range(0, count, n)]
    elif mode == "ranges":
        groups = parse_ranges(p.get("ranges", ""), count)
    else:
        groups = [[i] for i in range(count)]
    if len(groups) < 2 and mode != "ranges":
        raise ConversionError("This PDF has only one page.")
    folder = ctx.output_dir(f"{src.stem} split")
    for n, group in enumerate(groups, 1):
        ctx.check()
        writer = PdfWriter()
        for i in group:
            writer.add_page(reader.pages[i])
        name = f"p{group[0] + 1}" if len(group) == 1 else f"p{group[0] + 1}-{group[-1] + 1}"
        with open(folder / f"{src.stem} {name}.pdf", "wb") as fh:
            writer.write(fh)
        ctx.progress(n / len(groups))
    return [folder]


def page_count(path: Path) -> int:
    return len(_reader(path).pages)


def organize(ctx):
    """pages: [[index, rotation], ...] in the new order; repeats duplicate a page."""
    src = ctx.inputs[0]
    reader = _reader(src)
    pages = ctx.params.get("pages") or [[i, 0] for i in reversed(range(len(reader.pages)))]
    if any(i < 0 or i >= len(reader.pages) for i, _ in pages):
        raise ConversionError("The page list doesn't match this PDF.")
    writer = PdfWriter()
    for i, rotation in pages:
        page = writer.add_page(reader.pages[i])
        if rotation % 360:
            page.rotate(rotation % 360)
    out = ctx.output(f"{src.stem} (organized)", "pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    return [out]


def is_locked(path: Path) -> bool:
    try:
        return PdfReader(str(path)).is_encrypted
    except Exception:
        return False


def password(ctx):
    """Lock with a password (AES-256), or unlock a PDF whose password you know."""
    src = ctx.inputs[0]
    secret = ctx.params.get("password", "")
    reader = PdfReader(str(src))
    if reader.is_encrypted:
        if not reader.decrypt(secret):
            raise ConversionError("Wrong password.")
        writer = PdfWriter(clone_from=reader)
        out = ctx.output(f"{src.stem} (unlocked)", "pdf")
    else:
        if not secret:
            raise ConversionError("Type a password first.")
        writer = PdfWriter(clone_from=reader)
        writer.encrypt(user_password=secret, owner_password=None, algorithm="AES-256")
        out = ctx.output(f"{src.stem} (locked)", "pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    return [out]


def compress(ctx):
    src = ctx.inputs[0]
    _reader(src)  # friendly error for locked files
    writer = PdfWriter(clone_from=str(src))
    pages = list(writer.pages)
    for i, page in enumerate(pages):
        ctx.check()
        for image in page.images:
            try:
                picture = image.image
                if picture.width * picture.height < 250_000:
                    continue
                if max(picture.size) > 2200:
                    picture.thumbnail((2200, 2200))
                if picture.mode not in ("RGB", "L"):
                    picture = picture.convert("RGB")
                image.replace(picture, quality=65)
            except Exception:
                continue  # leave images we can't recompress as they are
        page.compress_content_streams()
        ctx.progress(0.9 * (i + 1) / len(pages))
    writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    out = ctx.output(f"{src.stem} (compressed)", "pdf")
    with open(out, "wb") as fh:
        writer.write(fh)
    if out.stat().st_size >= src.stat().st_size:
        raise ConversionError("Already compact: a smaller copy wasn't possible.")
    return [out]


def thumbnails(path: Path, height: int = 180):
    """Small page images for the reorder dialog."""
    pdf = _open_pdfium(path)
    try:
        images = []
        for i in range(len(pdf)):
            page = pdf[i]
            _, h_pt = page.get_size()
            images.append(page.render(scale=height / h_pt).to_pil())
        return images
    finally:
        pdf.close()
