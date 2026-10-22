"""Spreadsheets: XLSX with openpyxl, CSV and TSV with the csv module,
XLS and PDF output through Excel."""

from __future__ import annotations

import csv
import datetime as dt
import re
import tempfile
from pathlib import Path

from spinvert.engines import office
from spinvert.util import ConversionError, ext_of, safe_name

INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
FLOAT = re.compile(r"^-?(0|[1-9]\d*)?\.\d+$")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _read_delimited(path: Path, delimiter: str | None):
    text = _read_text(path)
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:20000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [row for row in csv.reader(text.splitlines(), delimiter=delimiter)]


def _cell_out(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _cell_in(text: str):
    if INT.match(text):
        return int(text)
    if FLOAT.match(text):
        return float(text)
    return text


def read_sheets(ctx, path: Path) -> list[tuple[str, list[list]]]:
    kind = ext_of(path)
    if kind == "csv":
        return [(path.stem, _read_delimited(path, None))]
    if kind == "tsv":
        return [(path.stem, _read_delimited(path, "\t"))]
    if kind == "xls":
        with tempfile.TemporaryDirectory(prefix="spinvert-") as tmp:
            xlsx = Path(tmp) / f"{path.stem}.xlsx"
            office.excel_save(path, xlsx, "xlsx")
            return read_sheets(ctx, xlsx)
    from openpyxl import load_workbook

    try:
        book = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:
        raise ConversionError(f"Can't read {path.name}: {e}") from e
    sheets = []
    for ws in book.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        while rows and all(v is None for v in rows[-1]):
            rows.pop()
        sheets.append((ws.title, rows))
    book.close()
    return sheets


def write_xlsx(sheets, out: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        ws = book.create_sheet(safe_name(name)[:31])
        for row in rows:
            ws.append([_cell_in(v) if isinstance(v, str) else v for v in row])
    if not book.worksheets:
        book.create_sheet("Sheet1")
    book.save(str(out))


def to_delimited(ext: str):
    delimiter = "\t" if ext == "tsv" else ","

    def run(ctx):
        src = ctx.inputs[0]
        sheets = [s for s in read_sheets(ctx, src) if s[1]] or [(src.stem, [])]
        outs = []
        for name, rows in sheets:
            stem = src.stem if len(sheets) == 1 else f"{src.stem} - {safe_name(name)}"
            out = ctx.output(stem, ext)
            with open(out, "w", encoding="utf-8-sig", newline="") as fh:
                writer = csv.writer(fh, delimiter=delimiter)
                for row in rows:
                    writer.writerow([_cell_out(v) for v in row])
            outs.append(out)
        return outs
    return run


def to_xlsx(ctx):
    src = ctx.inputs[0]
    out = ctx.output(src.stem, "xlsx")
    if ext_of(src) == "xls":
        office.excel_save(src, out, "xlsx")
    else:
        write_xlsx(read_sheets(ctx, src), out)
    return [out]


def to_pdf(ctx):
    src = ctx.inputs[0]
    out = ctx.output(src.stem, "pdf")
    if ext_of(src) in ("xlsx", "xls"):
        office.excel_save(src, out, "pdf")
        return [out]
    with tempfile.TemporaryDirectory(prefix="spinvert-") as tmp:
        xlsx = Path(tmp) / f"{src.stem}.xlsx"
        write_xlsx(read_sheets(ctx, src), xlsx)
        office.excel_save(xlsx, out, "pdf")
    return [out]
