"""Slides. PowerPoint renders PDF and images; python-pptx reads the text."""

from __future__ import annotations

from spinvert.engines import office
from spinvert.util import ConversionError


def slides_to_pdf(ctx):
    src = ctx.inputs[0]
    out = ctx.output(src.stem, "pdf")
    office.powerpoint_export(src, "pdf", out=out)
    return [out]


def slides_to_images(fmt: str):
    def run(ctx):
        src = ctx.inputs[0]
        folder = ctx.output_dir(f"{src.stem} slides")
        office.powerpoint_export(src, fmt, folder=folder, progress=ctx.progress)
        return [folder]
    return run


def _slide_texts(src):
    from pptx import Presentation

    try:
        prs = Presentation(str(src))
    except Exception as e:
        raise ConversionError(f"Can't read {src.name}: {e}") from e
    slides = []
    for slide in prs.slides:
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append((para.level, text))
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    lines.append((0, " | ".join(cells)))
        slides.append((title, lines))
    return slides


def slides_to_text(fmt: str):
    def run(ctx):
        src = ctx.inputs[0]
        out = ctx.output(src.stem, fmt)
        chunks = []
        for n, (title, lines) in enumerate(_slide_texts(src), 1):
            if fmt == "md":
                chunk = [f"## {title or f'Slide {n}'}", ""]
                chunk += [f"{'  ' * level}- {text}" for level, text in lines]
            else:
                chunk = [title or f"Slide {n}"]
                chunk += [f"{'  ' * level}{text}" for level, text in lines]
            chunks.append("\n".join(chunk).rstrip())
        out.write_text("\n\n".join(chunks) + "\n", encoding="utf-8")
        return [out]
    return run
